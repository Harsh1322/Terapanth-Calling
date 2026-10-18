import csv
import io
import json
import math
import re
from datetime import date
from functools import wraps
from zipfile import BadZipFile

from flask import (
    Blueprint,
    Response,
    abort,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    url_for,
)
from flask_login import current_user, login_required, login_user, logout_user
from flask_wtf.csrf import generate_csrf
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, select

from .extensions import db
from .models import AuditLog, Contact, DataCollection, DataTask, DataTaskAssignment, User

tracker = Blueprint("tracker", __name__)
KILPAUK_AREAS = {
    "ayanavaram",
    "kilpauk",
    "vepery",
    "puraswalkam",
    "purasawalkam",
    "northtown",
    "icf",
}
TASK_FIELD_TYPES = {"text", "long_text", "number", "date", "boolean"}


def classify_contact(age, area):
    if age > 21:
        return "typ"
    normalized_area = re.sub(r"[^a-z0-9]", "", area.lower())
    if normalized_area in KILPAUK_AREAS:
        return "kilpauk"
    return "success"


def validate_task_fields(fields):
    if not isinstance(fields, list) or len(fields) > 30:
        raise ValueError("Add no more than 30 data fields to a task.")
    clean_fields = []
    used_keys = set()
    for field in fields:
        if not isinstance(field, dict):
            raise ValueError("Each task field must include a name and type.")
        label = str(field.get("label", "")).strip()
        key = str(field.get("key", "")).strip()
        field_type = field.get("type", "text")
        required = field.get("required", True)
        if (
            not label
            or len(label) > 80
            or not re.fullmatch(r"[a-z][a-z0-9_]{0,39}", key)
            or key in {"age", "area"}
            or key in used_keys
            or not isinstance(field_type, str)
            or field_type not in TASK_FIELD_TYPES
            or not isinstance(required, bool)
        ):
            raise ValueError(
                "Check each field name, unique key, and type. Age and area are collected separately."
            )
        used_keys.add(key)
        clean_fields.append(
            {"key": key, "label": label, "type": field_type, "required": required}
        )
    return clean_fields


def task_payload(task):
    return {
        "id": task.id,
        "title": task.title,
        "instructions": task.instructions,
        "fields": json.loads(task.fields_json),
        "active": task.is_active,
        "assignedMembers": [
            {
                "id": assignment.member.id,
                "name": assignment.member.username,
                "active": assignment.member.is_active,
            }
            for assignment in task.assignments
        ],
    }


def has_active_data_task(member_id):
    return db.session.scalar(
        select(DataTaskAssignment.id)
        .join(DataTask)
        .where(
            DataTaskAssignment.member_id == member_id,
            DataTask.is_active.is_(True),
        )
        .limit(1)
    ) is not None


def contact_payload(contact, requires_collection=None):
    if requires_collection is None:
        requires_collection = bool(
            contact.assigned_to_id and has_active_data_task(contact.assigned_to_id)
        )
    return {
        "id": contact.id,
        "name": contact.name or "Unnamed contact",
        "phone": contact.phone,
        "status": contact.status,
        "statusDisplay": contact.status_display,
        "destination": contact.destination,
        "destinationDisplay": contact.destination_display if contact.destination else "",
        "assignedTo": contact.assigned_to.username if contact.assigned_to else "Unassigned",
        "requiresCollection": requires_collection,
    }


def admin_required(view):
    @wraps(view)
    @login_required
    def wrapped(*args, **kwargs):
        if not current_user.is_superuser:
            if request.path.startswith("/api/"):
                return jsonify({"error": "Administrator access is required."}), 403
            abort(403)
        return view(*args, **kwargs)

    return wrapped


@tracker.route("/login/", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("tracker.dashboard"))
    error = None
    mobile_value = ""
    if request.method == "POST":
        mobile_value = request.form.get("mobile", "")
        mobile = re.sub(r"\D", "", mobile_value)
        user = db.session.scalar(
            select(User).where(User.mobile == mobile, User.is_active.is_(True))
        )
        if user is None:
            error = "No active account uses that mobile number."
        else:
            login_user(user)
            return redirect(url_for("tracker.dashboard"))
    return render_template("login.html", error=error, mobile=mobile_value)


@tracker.route("/logout/")
@login_required
def logout():
    logout_user()
    return redirect(url_for("tracker.login"))


@tracker.get("/api/csrf")
def api_csrf():
    return jsonify({"csrfToken": generate_csrf()})


@tracker.get("/api/health")
def api_health():
    db.session.execute(select(1)).scalar_one()
    return jsonify({"status": "ok"})


@tracker.get("/api/session")
def api_session():
    if not current_user.is_authenticated:
        return jsonify({"user": None})
    return jsonify(
        {
            "user": {
                "id": current_user.id,
                "name": current_user.username,
                "mobile": current_user.mobile,
                "isAdmin": current_user.is_superuser,
            }
        }
    )


@tracker.post("/api/login")
def api_login():
    data = request.get_json(silent=True)
    mobile_value = data.get("mobile", "") if isinstance(data, dict) else ""
    mobile = re.sub(r"\D", "", str(mobile_value))
    user = db.session.scalar(
        select(User).where(User.mobile == mobile, User.is_active.is_(True))
    )
    if user is None:
        return jsonify({"error": "No active account uses that mobile number."}), 401
    login_user(user)
    return jsonify({"user": {"name": user.username, "isAdmin": user.is_superuser}})


@tracker.post("/api/logout")
@login_required
def api_logout():
    logout_user()
    return jsonify({"ok": True})


@tracker.get("/api/dashboard")
@login_required
def api_dashboard():
    stats_query = select(Contact.status, func.count(Contact.id)).group_by(Contact.status)
    contact_query = select(Contact).order_by(Contact.name, Contact.phone).limit(500)
    if not current_user.is_superuser:
        stats_query = stats_query.where(Contact.assigned_to_id == current_user.id)
        contact_query = contact_query.where(Contact.assigned_to_id == current_user.id)

    status_counts = dict(db.session.execute(stats_query).all())
    stats = {
        "total": sum(status_counts.values()),
        "pending": status_counts.get("pending", 0),
        "typ_member": status_counts.get("typ_member", 0),
        "kilpauk_member": status_counts.get("kilpauk_member", 0),
        "success": status_counts.get("success", 0),
        "reached": (
            status_counts.get("success", 0)
            + status_counts.get("typ_member", 0)
            + status_counts.get("kilpauk_member", 0)
        ),
        "no_answer": status_counts.get("no_answer", 0),
        "invalid": status_counts.get("invalid", 0),
        "follow_up": status_counts.get("follow_up", 0),
    }
    contacts = db.session.scalars(contact_query).all()
    members_with_tasks = set(
        db.session.scalars(
            select(DataTaskAssignment.member_id)
            .join(DataTask)
            .where(DataTask.is_active.is_(True))
        ).all()
    )
    result = {
        "stats": stats,
        "contacts": [
            contact_payload(contact, contact.assigned_to_id in members_with_tasks)
            for contact in contacts
        ],
        "user": {
            "id": current_user.id,
            "name": current_user.username,
            "mobile": current_user.mobile,
            "isAdmin": current_user.is_superuser,
        },
        "requiresCollection": (
            not current_user.is_superuser and has_active_data_task(current_user.id)
        ),
    }

    if current_user.is_superuser:
        members = db.session.scalars(
            select(User)
            .where(User.is_committee_member.is_(True), User.is_active.is_(True))
            .order_by(User.username)
        ).all()
        assignment_counts = {}
        for member_id, status, count in db.session.execute(
            select(Contact.assigned_to_id, Contact.status, func.count(Contact.id))
            .group_by(Contact.assigned_to_id, Contact.status)
        ).all():
            assignment_counts.setdefault(member_id, {})[status] = count
        result["memberStats"] = [
            {
                "id": member.id,
                "name": member.username,
                "mobile": member.mobile,
                "total": sum(counts.values()),
                "success": counts.get("success", 0),
                "reached": (
                    counts.get("success", 0)
                    + counts.get("typ_member", 0)
                    + counts.get("kilpauk_member", 0)
                ),
                "pending": sum(
                    counts.get(status, 0)
                    for status in ("pending", "follow_up", "no_answer")
                ),
                "invalid": counts.get("invalid", 0),
            }
            for member in members
            for counts in [assignment_counts.get(member.id, {})]
        ]
    return jsonify(result)


@tracker.get("/api/tasks")
@login_required
def api_tasks():
    query = select(DataTask).order_by(DataTask.created_at.desc(), DataTask.id.desc())
    if not current_user.is_superuser:
        query = (
            query.join(DataTaskAssignment)
            .where(
                DataTaskAssignment.member_id == current_user.id,
                DataTask.is_active.is_(True),
            )
            .distinct()
        )
    tasks = db.session.scalars(query).all()
    return jsonify({"tasks": [task_payload(task) for task in tasks]})


def task_form_data(data):
    if not isinstance(data, dict):
        raise ValueError("A JSON request body is required.")
    title = str(data.get("title", "")).strip()
    instructions = str(data.get("instructions", "")).strip()
    if not title or len(title) > 150 or not instructions or len(instructions) > 4000:
        raise ValueError("Enter a title and instructions (up to 4,000 characters).")

    fields = validate_task_fields(data.get("fields"))
    member_ids = data.get("memberIds")
    if (
        not isinstance(member_ids, list)
        or not member_ids
        or any(
            not isinstance(member_id, int) or isinstance(member_id, bool)
            for member_id in member_ids
        )
    ):
        raise ValueError("Assign this task to at least one active committee member.")
    member_ids = list(dict.fromkeys(member_ids))
    members = db.session.scalars(
        select(User).where(
            User.id.in_(member_ids),
            User.is_committee_member.is_(True),
            User.is_active.is_(True),
        )
    ).all()
    if len(members) != len(member_ids):
        raise ValueError("Tasks can only be assigned to active committee members.")
    return title, instructions, fields, members


def replace_task_assignments(task, members):
    task.assignments.clear()
    db.session.flush()
    task.assignments = [DataTaskAssignment(member=member) for member in members]


@tracker.post("/api/tasks")
@admin_required
def api_create_task():
    try:
        title, instructions, fields, members = task_form_data(request.get_json(silent=True))
    except ValueError as error:
        return jsonify({"error": str(error)}), 400

    task = DataTask(
        title=title,
        instructions=instructions,
        fields_json=json.dumps(fields),
        created_by=current_user,
    )
    task.assignments = [DataTaskAssignment(member=member) for member in members]
    db.session.add(task)
    db.session.flush()
    db.session.add(
        AuditLog(
            actor=current_user,
            action="data_task_created",
            details=f"Created task {task.id}: {task.title}",
        )
    )
    db.session.commit()
    return jsonify({"task": task_payload(task)}), 201


@tracker.patch("/api/tasks/<int:task_id>")
@admin_required
def api_update_task(task_id):
    task = db.session.get(DataTask, task_id)
    if task is None:
        return jsonify({"error": "Task not found."}), 404
    try:
        title, instructions, fields, members = task_form_data(request.get_json(silent=True))
    except ValueError as error:
        return jsonify({"error": str(error)}), 400

    task.title = title
    task.instructions = instructions
    task.fields_json = json.dumps(fields)
    replace_task_assignments(task, members)
    db.session.add(
        AuditLog(
            actor=current_user,
            action="data_task_updated",
            details=f"Updated task {task.id}: {task.title}",
        )
    )
    db.session.commit()
    return jsonify({"task": task_payload(task)})


@tracker.delete("/api/tasks/<int:task_id>")
@admin_required
def api_disable_task(task_id):
    task = db.session.get(DataTask, task_id)
    if task is None:
        return jsonify({"error": "Task not found."}), 404
    task.is_active = False
    db.session.add(
        AuditLog(
            actor=current_user,
            action="data_task_disabled",
            details=f"Disabled task {task.id}: {task.title}",
        )
    )
    db.session.commit()
    return jsonify({"ok": True})


@tracker.post("/api/tasks/<int:task_id>/enable")
@admin_required
def api_enable_task(task_id):
    task = db.session.get(DataTask, task_id)
    if task is None:
        return jsonify({"error": "Task not found."}), 404
    if not any(assignment.member.is_active for assignment in task.assignments):
        return jsonify({"error": "Assign the task to an active member before enabling it."}), 400
    task.is_active = True
    db.session.add(
        AuditLog(
            actor=current_user,
            action="data_task_enabled",
            details=f"Enabled task {task.id}: {task.title}",
        )
    )
    db.session.commit()
    return jsonify({"task": task_payload(task)})


@tracker.get("/api/tasks/<int:task_id>/contacts")
@login_required
def api_task_contacts(task_id):
    task = db.session.get(DataTask, task_id)
    if task is None or not task.is_active:
        return jsonify({"error": "This task is no longer active."}), 404
    assigned = db.session.scalar(
        select(DataTaskAssignment.id).where(
            DataTaskAssignment.task_id == task_id,
            DataTaskAssignment.member_id == current_user.id,
        )
    )
    if current_user.is_superuser or assigned is None:
        return jsonify({"error": "This task is not assigned to your account."}), 403

    contacts = db.session.scalars(
        select(Contact)
        .where(Contact.assigned_to_id == current_user.id)
        .order_by(Contact.name, Contact.phone)
    ).all()
    latest_by_contact = {}
    for submission in db.session.scalars(
        select(DataCollection)
        .where(
            DataCollection.task_id == task_id,
            DataCollection.member_id == current_user.id,
        )
        .order_by(DataCollection.submitted_at.desc(), DataCollection.id.desc())
    ).all():
        latest_by_contact.setdefault(submission.contact_id, submission)

    return jsonify(
        {
            "contacts": [
                {
                    **contact_payload(contact, True),
                    "submission": (
                        {
                            "age": latest_by_contact[contact.id].age,
                            "area": latest_by_contact[contact.id].area,
                            "classification": latest_by_contact[contact.id].classification,
                            "fields": json.loads(latest_by_contact[contact.id].fields_json),
                        }
                        if contact.id in latest_by_contact
                        else None
                    ),
                }
                for contact in contacts
            ]
        }
    )


def validate_task_answers(fields, answers):
    if not isinstance(answers, dict):
        raise ValueError("Enter the task’s required information.")
    cleaned = {}
    for field in fields:
        value = answers.get(field["key"])
        if value is None:
            value = ""
        if field["type"] == "boolean":
            if not isinstance(value, bool):
                raise ValueError(f"{field['label']} must be answered yes or no.")
            cleaned[field["key"]] = value
            continue
        if not isinstance(value, str):
            raise ValueError(f"Enter a valid value for {field['label']}.")
        value = value.strip()
        if field["required"] and not value:
            raise ValueError(f"{field['label']} is required for this task.")
        if len(value) > 2000:
            raise ValueError(f"{field['label']} must be 2,000 characters or fewer.")
        if value and field["type"] == "number":
            try:
                parsed = float(value)
            except ValueError as error:
                raise ValueError(f"Enter a number for {field['label']}.") from error
            if not math.isfinite(parsed):
                raise ValueError(f"Enter a number for {field['label']}.")
        if value and field["type"] == "date":
            try:
                date.fromisoformat(value)
            except ValueError as error:
                raise ValueError(f"Enter a valid date for {field['label']}.") from error
        cleaned[field["key"]] = value
    return cleaned


@tracker.post("/api/tasks/<int:task_id>/submissions")
@login_required
def api_submit_task(task_id):
    if current_user.is_superuser:
        return jsonify({"error": "Only the assigned committee member can submit this task."}), 403
    task = db.session.get(DataTask, task_id)
    if task is None or not task.is_active:
        return jsonify({"error": "This task is no longer active."}), 404
    assigned = db.session.scalar(
        select(DataTaskAssignment.id).where(
            DataTaskAssignment.task_id == task_id,
            DataTaskAssignment.member_id == current_user.id,
        )
    )
    if assigned is None:
        return jsonify({"error": "This task is not assigned to your account."}), 403

    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"error": "A JSON request body is required."}), 400
    contact_id = data.get("contactId")
    age = data.get("age")
    area = data.get("area")
    if not isinstance(contact_id, int):
        return jsonify({"error": "Choose a contact assigned to you."}), 400
    contact = db.session.scalar(
        select(Contact).where(
            Contact.id == contact_id,
            Contact.assigned_to_id == current_user.id,
        )
    )
    if contact is None:
        return jsonify({"error": "This contact is not assigned to your account."}), 403
    if isinstance(age, bool) or not isinstance(age, int) or age < 0 or age > 130:
        return jsonify({"error": "Enter an age between 0 and 130."}), 400
    if not isinstance(area, str) or not area.strip() or len(area.strip()) > 120:
        return jsonify({"error": "Choose or enter the contact’s area."}), 400

    fields = json.loads(task.fields_json)
    try:
        answers = validate_task_answers(fields, data.get("fields", {}))
    except ValueError as error:
        return jsonify({"error": str(error)}), 400

    classification = classify_contact(age, area.strip())
    status, destination = {
        "typ": ("typ_member", "typ"),
        "kilpauk": ("kilpauk_member", "kilpauk"),
        "success": ("success", "other"),
    }[classification]
    submission = DataCollection(
        task=task,
        contact=contact,
        member=current_user,
        task_title=task.title,
        age=age,
        area=area.strip(),
        classification=classification,
        fields_json=json.dumps(answers),
        field_labels_json=json.dumps(
            {field["key"]: field["label"] for field in fields}
        ),
    )
    contact.status = status
    contact.destination = destination
    db.session.add(submission)
    db.session.add(
        AuditLog(
            actor=current_user,
            contact=contact,
            action="data_collection_submitted",
            details=f"Task {task.id}; classification={classification}",
        )
    )
    db.session.commit()
    return jsonify(
        {
            "ok": True,
            "status": status,
            "classification": classification,
            "classificationDisplay": {
                "typ": "TYP member",
                "kilpauk": "Kilpauk member",
                "success": "Success — details collected",
            }[classification],
        }
    )


@tracker.route("/")
@login_required
def dashboard():
    stats_query = select(Contact.status, func.count(Contact.id)).group_by(Contact.status)
    contact_query = select(Contact).order_by(Contact.name, Contact.phone).limit(500)
    if not current_user.is_superuser:
        stats_query = stats_query.where(Contact.assigned_to_id == current_user.id)
        contact_query = contact_query.where(Contact.assigned_to_id == current_user.id)
    status_counts = dict(db.session.execute(stats_query).all())
    stats = {
        "total": sum(status_counts.values()),
        "pending": status_counts.get("pending", 0),
        "success": status_counts.get("success", 0),
        "no_answer": status_counts.get("no_answer", 0),
        "invalid": status_counts.get("invalid", 0),
        "follow_up": status_counts.get("follow_up", 0),
    }
    contacts = db.session.scalars(contact_query).all()
    members = db.session.scalars(
        select(User)
        .where(User.is_committee_member.is_(True), User.is_active.is_(True))
        .order_by(User.username)
    ).all()
    member_stats = []
    if current_user.is_superuser:
        assignment_counts = {}
        results = db.session.execute(
            select(Contact.assigned_to_id, Contact.status, func.count(Contact.id))
            .group_by(Contact.assigned_to_id, Contact.status)
        ).all()
        for member_id, status, count in results:
            assignment_counts.setdefault(member_id, {})[status] = count
        for member in members:
            counts = assignment_counts.get(member.id, {})
            member_stats.append(
                {
                    "member": member,
                    "total": sum(counts.values()),
                    "success": (
                        counts.get("success", 0)
                        + counts.get("typ_member", 0)
                        + counts.get("kilpauk_member", 0)
                    ),
                    "pending": sum(
                        counts.get(status, 0)
                        for status in ("pending", "follow_up", "no_answer")
                    ),
                    "invalid": counts.get("invalid", 0),
                }
            )
    return render_template(
        "dashboard.html",
        contacts=contacts,
        stats=stats,
        member_stats=member_stats,
    )


@tracker.post("/api/contacts/<int:contact_id>/status")
@login_required
def api_update_status(contact_id):
    contact = db.session.get(Contact, contact_id)
    if contact is None:
        return jsonify({"error": "Contact not found."}), 404
    if not current_user.is_superuser and contact.assigned_to_id != current_user.id:
        return jsonify({"error": "You cannot update this contact."}), 403

    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"error": "A JSON request body is required."}), 400
    status = data.get("status")
    destination = data.get("destination", "")
    allowed_statuses = dict(Contact.STATUS_CHOICES)
    allowed_destinations = dict(Contact.DESTINATION_CHOICES)
    if not isinstance(status, str) or status not in allowed_statuses:
        return jsonify({"error": "Choose a valid contact status."}), 400
    if (
        status in ("success", "typ_member", "kilpauk_member")
        and contact.assigned_to_id
        and has_active_data_task(contact.assigned_to_id)
    ):
        return jsonify(
            {"error": "Complete an assigned data task to record a successful call outcome."}
        ), 400
    if status == "typ_member":
        destination = "typ"
    elif status == "kilpauk_member":
        destination = "kilpauk"
    elif status == "success" and (
        not isinstance(destination, str) or destination not in allowed_destinations
    ):
        return jsonify(
            {"error": "Choose a destination for a successful contact."}
        ), 400
    if status not in ("success", "typ_member", "kilpauk_member"):
        destination = ""

    previous = f"{contact.status}/{contact.destination}"
    contact.status = status
    contact.destination = destination
    contact.notes = str(data.get("notes", ""))[:500]
    db.session.add(
        AuditLog(
            actor=current_user,
            contact=contact,
            action="status_updated",
            details=f"{previous} -> {status}/{destination}",
        )
    )
    db.session.commit()
    return jsonify({"contact": contact_payload(contact)})


@tracker.route("/members/", methods=["GET", "POST"])
@admin_required
def members():
    if request.method == "POST":
        mobile = re.sub(r"\D", "", request.form.get("mobile", ""))
        display_name = request.form.get("display_name", "").strip()
        password = request.form.get("password", "")
        if (
            not mobile
            or len(mobile) > 20
            or not display_name
            or len(display_name) > 150
        ):
            flash("Enter a valid mobile number and a name of at most 150 characters.", "error")
        elif len(password) < 8:
            flash("The password must contain at least 8 characters.", "error")
        elif db.session.scalar(select(User.id).where(User.mobile == mobile)) is not None:
            flash("A user with that mobile number already exists.", "error")
        else:
            user = User(
                username=display_name,
                mobile=mobile,
                is_committee_member=True,
            )
            user.set_password(password)
            db.session.add(user)
            db.session.add(
                AuditLog(
                    actor=current_user,
                    action="member_added",
                    details=f"Added member {mobile}",
                )
            )
            db.session.commit()
            flash(f"Member {mobile} created.", "success")
            return redirect(url_for("tracker.members"))
    member_list = db.session.scalars(
        select(User)
        .where(User.is_committee_member.is_(True))
        .order_by(User.username)
    ).all()
    return render_template("members.html", members=member_list)


@tracker.get("/api/members")
@admin_required
def api_members():
    member_list = db.session.scalars(
        select(User)
        .where(User.is_committee_member.is_(True))
        .order_by(User.username)
    ).all()
    return jsonify(
        {
            "members": [
                {
                    "id": member.id,
                    "name": member.username,
                    "mobile": member.mobile,
                    "active": member.is_active,
                }
                for member in member_list
            ]
        }
    )


@tracker.post("/api/members")
@admin_required
def api_add_member():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"error": "A JSON request body is required."}), 400
    mobile = re.sub(r"\D", "", str(data.get("mobile", "")))
    display_name = str(data.get("displayName", "")).strip()
    password = str(data.get("password", ""))
    if (
        not mobile
        or len(mobile) > 20
        or not display_name
        or len(display_name) > 150
    ):
        return jsonify(
            {"error": "Enter a valid mobile number and a name of at most 150 characters."}
        ), 400
    if len(password) < 8:
        return jsonify({"error": "The password must contain at least 8 characters."}), 400
    if db.session.scalar(select(User.id).where(User.mobile == mobile)) is not None:
        return jsonify({"error": "A user with that mobile number already exists."}), 409

    user = User(
        username=display_name,
        mobile=mobile,
        is_committee_member=True,
    )
    user.set_password(password)
    db.session.add(user)
    db.session.add(
        AuditLog(
            actor=current_user,
            action="member_added",
            details=f"Added member {mobile}",
        )
    )
    db.session.commit()
    return jsonify(
        {
            "member": {
                "id": user.id,
                "name": user.username,
                "mobile": user.mobile,
                "active": user.is_active,
            }
        }
    ), 201


@tracker.patch("/api/members/<int:user_id>")
@admin_required
def api_update_member(user_id):
    user = db.session.scalar(
        select(User).where(
            User.id == user_id,
            User.is_committee_member.is_(True),
        )
    )
    if user is None:
        return jsonify({"error": "Member not found."}), 404
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"error": "A JSON request body is required."}), 400

    display_name = str(data.get("displayName", "")).strip()
    mobile = re.sub(r"\D", "", str(data.get("mobile", "")))
    if (
        not display_name
        or len(display_name) > 150
        or not mobile
        or len(mobile) > 20
    ):
        return jsonify(
            {"error": "Enter a name of at most 150 characters and a valid mobile number."}
        ), 400
    duplicate = db.session.scalar(
        select(User.id).where(User.mobile == mobile, User.id != user.id)
    )
    if duplicate is not None:
        return jsonify({"error": "A user with that mobile number already exists."}), 409

    previous = f"{user.username}/{user.mobile}"
    user.username = display_name
    user.mobile = mobile
    db.session.add(
        AuditLog(
            actor=current_user,
            action="member_updated",
            details=f"Updated member {user.id}: {previous} -> {display_name}/{mobile}",
        )
    )
    db.session.commit()
    return jsonify(
        {
            "member": {
                "id": user.id,
                "name": user.username,
                "mobile": user.mobile,
                "active": user.is_active,
            }
        }
    )


@tracker.route("/members/<int:user_id>/delete/", methods=["POST"])
@admin_required
def delete_member(user_id):
    user = db.session.scalar(
        select(User).where(
            User.id == user_id,
            User.is_committee_member.is_(True),
        )
    )
    if user is None:
        abort(404)
    mobile = user.mobile
    user.is_active = False
    db.session.add(
        AuditLog(
            actor=current_user,
            action="member_deactivated",
            details=f"Deactivated {mobile}; contact ownership retained",
        )
    )
    db.session.commit()
    flash("Member deactivated. Existing assignments were retained for audit.", "success")
    return redirect(url_for("tracker.members"))


@tracker.post("/api/members/<int:user_id>/deactivate")
@admin_required
def api_deactivate_member(user_id):
    user = db.session.scalar(
        select(User).where(
            User.id == user_id,
            User.is_committee_member.is_(True),
        )
    )
    if user is None:
        return jsonify({"error": "Member not found."}), 404
    user.is_active = False
    db.session.add(
        AuditLog(
            actor=current_user,
            action="member_deactivated",
            details=f"Deactivated {user.mobile}; contact ownership retained",
        )
    )
    db.session.commit()
    return jsonify({"ok": True})


def normalize_phone(raw):
    value = re.sub(r"[^0-9+]", "", str(raw or ""))
    digits = re.sub(r"\D", "", value)
    if len(digits) == 10:
        return "+91" + digits
    if len(digits) == 12 and digits.startswith("91"):
        return "+" + digits
    if value.startswith("+") and len(digits) >= 8:
        return "+" + digits
    return digits


def read_rows(upload):
    filename = upload.filename.lower()
    if filename.endswith(".csv"):
        content = upload.read().decode("utf-8-sig", errors="replace")
        return list(csv.DictReader(io.StringIO(content)))
    if filename.endswith(".xlsx"):
        workbook = load_workbook(upload, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = [str(value or "").strip().lower() for value in next(rows, [])]
        return [dict(zip(headers, row)) for row in rows]
    raise ValueError("Upload a .csv or .xlsx file.")


def field(row, choices):
    normalized = {
        re.sub(r"[^a-z0-9]", "", str(key).lower()): value
        for key, value in row.items()
    }
    for key in choices:
        if key in normalized and normalized[key] is not None:
            return str(normalized[key]).strip()
    return ""


def import_contacts_from_upload(upload):
    try:
        rows = read_rows(upload)
    except (ValueError, BadZipFile, InvalidFileException, OSError) as error:
        raise ValueError(f"Could not read file: {error}") from error

    cleaned = []
    seen = set()
    for row in rows:
        phone = normalize_phone(
            field(row, ["phone", "mobile", "mobilenumber", "phonenumber", "number", "telephone"])
        )
        if not phone or phone in seen:
            continue
        seen.add(phone)
        cleaned.append(
            {
                "phone": phone,
                "name": field(row, ["name", "fullname", "contactname", "displayname"]),
                "group": field(row, ["group", "groupname", "sourcegroup"]),
            }
        )

    existing = set(
        db.session.scalars(
            select(Contact.phone).where(Contact.phone.in_([item["phone"] for item in cleaned]))
        ).all()
    )
    new_rows = [item for item in cleaned if item["phone"] not in existing]
    member_list = db.session.scalars(
        select(User)
        .where(User.is_committee_member.is_(True), User.is_active.is_(True))
        .order_by(User.id)
    ).all()
    if not member_list and new_rows:
        raise ValueError("Add at least one active committee member before importing contacts.")

    for index, item in enumerate(new_rows):
        member = member_list[index % len(member_list)]
        db.session.add(
            Contact(
                name=item["name"][:180],
                phone=item["phone"],
                source_group=item["group"][:180],
                assigned_to=member,
            )
        )
    skipped = len(rows) - len(new_rows)
    db.session.add(
        AuditLog(
            actor=current_user,
            action="contacts_imported",
            details=(
                f"Rows read={len(rows)}, unique valid={len(cleaned)}, "
                f"created={len(new_rows)}, duplicates skipped={skipped}"
            ),
        )
    )
    db.session.commit()
    return len(new_rows), skipped


@tracker.route("/contacts/import/", methods=["POST"])
@admin_required
def import_contacts():
    upload = request.files.get("file")
    if upload is None or not upload.filename:
        return Response("Upload a CSV or XLSX file.", status=400)
    try:
        created, skipped = import_contacts_from_upload(upload)
    except ValueError as error:
        flash(str(error), "error")
        return redirect(url_for("tracker.dashboard"))
    flash(
        f"Imported {created} new contacts. "
        f"{skipped} duplicates/invalid or already-existing entries were skipped.",
        "success",
    )
    return redirect(url_for("tracker.dashboard"))


@tracker.post("/api/contacts/import")
@admin_required
def api_import_contacts():
    upload = request.files.get("file")
    if upload is None or not upload.filename:
        return jsonify({"error": "Upload a CSV or XLSX file."}), 400
    try:
        created, skipped = import_contacts_from_upload(upload)
    except ValueError as error:
        return jsonify({"error": str(error)}), 400
    return jsonify(
        {
            "created": created,
            "skipped": skipped,
            "message": (
                f"Imported {created} new contacts. "
                f"{skipped} duplicates, invalid entries, or existing contacts were skipped."
            ),
        }
    )


@tracker.route("/contacts/<int:contact_id>/status/", methods=["POST"])
@login_required
def update_status(contact_id):
    contact = db.session.get(Contact, contact_id)
    if contact is None:
        abort(404)
    if not current_user.is_superuser and contact.assigned_to_id != current_user.id:
        abort(403)
    status = request.form.get("status")
    destination = request.form.get("destination", "")
    allowed_statuses = dict(Contact.STATUS_CHOICES)
    allowed_destinations = dict(Contact.DESTINATION_CHOICES)
    if status not in allowed_statuses:
        return Response("Invalid status", status=400)
    if (
        status in ("success", "typ_member", "kilpauk_member")
        and contact.assigned_to_id
        and has_active_data_task(contact.assigned_to_id)
    ):
        return Response(
            "Complete an assigned data task to record a successful call outcome.",
            status=400,
        )
    if status == "typ_member":
        destination = "typ"
    elif status == "kilpauk_member":
        destination = "kilpauk"
    elif status == "success" and destination not in allowed_destinations:
        return Response(
            "Choose a destination for a successful contact.",
            status=400,
        )
    if status not in ("success", "typ_member", "kilpauk_member"):
        destination = ""
    previous = f"{contact.status}/{contact.destination}"
    contact.status = status
    contact.destination = destination
    contact.notes = request.form.get("notes", "")[:500]
    db.session.add(
        AuditLog(
            actor=current_user,
            contact=contact,
            action="status_updated",
            details=f"{previous} -> {status}/{destination}",
        )
    )
    db.session.commit()
    flash("Contact status updated.", "success")
    return redirect(url_for("tracker.dashboard"))


@tracker.route("/contacts/export.xlsx")
@tracker.route("/api/contacts/export.xlsx")
@admin_required
def export_excel():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "All Contacts"
    worksheet.append(
        [
            "Name",
            "Mobile",
            "Assigned Member",
            "Member Mobile",
            "Status",
            "Success Destination",
            "Source Group",
            "Created At",
            "Last Updated",
        ]
    )
    contacts = db.session.scalars(select(Contact).order_by(Contact.phone)).all()
    for contact in contacts:
        worksheet.append(
            [
                contact.name,
                contact.phone,
                contact.assigned_to.username if contact.assigned_to else "",
                contact.assigned_to.mobile if contact.assigned_to else "",
                contact.status_display,
                contact.destination_display if contact.destination else "",
                contact.source_group,
                contact.created_at.strftime("%Y-%m-%d %H:%M") if contact.created_at else "",
                contact.updated_at.strftime("%Y-%m-%d %H:%M") if contact.updated_at else "",
            ]
        )
    for column in worksheet.columns:
        letter = column[0].column_letter
        worksheet.column_dimensions[letter].width = min(
            max(max(len(str(cell.value or "")) for cell in column) + 2, 12),
            35,
        )
    audit_sheet = workbook.create_sheet("Audit Log")
    audit_sheet.append(["Timestamp", "Actor", "Action", "Contact Mobile", "Details"])
    logs = db.session.scalars(
        select(AuditLog).order_by(AuditLog.created_at.desc())
    ).all()
    for log in logs:
        audit_sheet.append(
            [
                log.created_at.strftime("%Y-%m-%d %H:%M"),
                log.actor.mobile if log.actor else "",
                log.action,
                log.contact.phone if log.contact else "",
                log.details,
            ]
        )
    collection_sheet = workbook.create_sheet("Data Collection")
    submissions = db.session.scalars(
        select(DataCollection).order_by(DataCollection.submitted_at, DataCollection.id)
    ).all()
    extra_fields = {}
    for submission in submissions:
        labels = json.loads(submission.field_labels_json)
        for key, label in labels.items():
            extra_fields[(submission.task_title, key)] = label
    base_headers = [
        "Task",
        "Contact",
        "Mobile",
        "Committee Member",
        "Age",
        "Area",
        "Classification",
        "Submitted At",
    ]
    extra_columns = sorted(extra_fields)
    collection_sheet.append(
        base_headers
        + [f"{title} — {extra_fields[(title, key)]}" for title, key in extra_columns]
    )
    for submission in submissions:
        answers = json.loads(submission.fields_json)
        collection_sheet.append(
            [
                submission.task_title,
                submission.contact.name,
                submission.contact.phone,
                submission.member.username,
                submission.age,
                submission.area,
                {
                    "typ": "TYP member",
                    "kilpauk": "Kilpauk member",
                    "success": "Success",
                }.get(submission.classification, submission.classification),
                submission.submitted_at.strftime("%Y-%m-%d %H:%M"),
            ]
            + [
                answers.get(key, "")
                if submission.task_title == title
                else ""
                for title, key in extra_columns
            ]
        )
    for column in collection_sheet.columns:
        letter = column[0].column_letter
        collection_sheet.column_dimensions[letter].width = min(
            max(max(len(str(cell.value or "")) for cell in column) + 2, 12),
            35,
        )
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="committee_contacts.xlsx",
    )


@tracker.route("/extension/import/", methods=["POST"])
@admin_required
def extension_import():
    try:
        data = json.loads(request.get_data(as_text=True))
        contacts = data.get("contacts", [])
        source = str(data.get("groupName", "WhatsApp group"))[:180]
    except (json.JSONDecodeError, AttributeError, TypeError):
        return Response("Invalid JSON", status=400)
    if not isinstance(contacts, list):
        return Response("Invalid contacts list", status=400)

    member_list = db.session.scalars(
        select(User)
        .where(User.is_committee_member.is_(True), User.is_active.is_(True))
        .order_by(User.id)
    ).all()
    if contacts and not member_list:
        return Response("Add active committee members first.", status=400)
    cleaned = []
    seen = set()
    for item in contacts:
        if not isinstance(item, dict):
            continue
        phone = normalize_phone(item.get("phone", ""))
        if phone and phone not in seen:
            exists = db.session.scalar(select(Contact.id).where(Contact.phone == phone))
            if exists is None:
                seen.add(phone)
                cleaned.append((str(item.get("name", ""))[:180], phone))
    for index, (name, phone) in enumerate(cleaned):
        db.session.add(
            Contact(
                name=name,
                phone=phone,
                assigned_to=member_list[index % len(member_list)],
                source_group=source,
            )
        )
    db.session.add(
        AuditLog(
            actor=current_user,
            action="whatsapp_extension_import",
            details=f"Group={source}; created={len(cleaned)}",
        )
    )
    db.session.commit()
    return jsonify(
        {"created": len(cleaned), "message": f"Imported {len(cleaned)} new contacts."}
    )
